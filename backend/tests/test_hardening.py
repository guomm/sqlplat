import pytest
from app.sql import validate_sql


def test_session_sql_mode_cannot_change_parser_rules():
    with pytest.raises(ValueError):
        validate_sql("SET SESSION sql_mode='NO_BACKSLASH_ESCAPES'; SELECT 1;")


def test_excel_control_characters_are_exportable(tmp_path):
    from app.results import ResultStore
    from openpyxl import load_workbook

    store = ResultStore(tmp_path)
    store.path("r.jsonl").write_text('["a\\u0001b"]\n')
    with store.xlsx_file("r.jsonl", [{"name": "c"}]) as file:
        assert load_workbook(file).active["A2"].value == "a�b"


@pytest.mark.parametrize(
    "sql",
    [
        'SET LOCAL sql_mode="NO_BACKSLASH_ESCAPES"',
        "SET @@local.autocommit=0",
        "/*! SET autocommit=0 */",
        '/*!40101 SET SESSION `sql_mode`="ANSI_QUOTES" */',
        "SET @a=1, autocommit=0",
    ],
)
def test_alternative_control_syntax_is_rejected(sql):
    with pytest.raises(ValueError):
        validate_sql(sql)


def test_excel_16_digit_integer_is_text_even_if_js_safe(tmp_path):
    from app.results import ResultStore
    from openpyxl import load_workbook

    store = ResultStore(tmp_path)
    store.path("r.jsonl").write_text("[1234567890123456]\n")
    with store.xlsx_file("r.jsonl", [{"name": "id"}]) as file:
        cell = load_workbook(file).active["A2"]
        assert cell.data_type == "s" and cell.value == "1234567890123456"


@pytest.mark.parametrize(
    "sql",
    [
        "SET @@autocommit=0",
        'SET @@sql_mode="NO_BACKSLASH_ESCAPES"',
        "SET @@SESSION.`autocommit`=0",
    ],
)
def test_unqualified_system_variables_are_rejected(sql):
    with pytest.raises(ValueError):
        validate_sql(sql)
